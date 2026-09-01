import csv
import io
import logging
import os
import re
import secrets
import time
import unicodedata
import uuid
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import List, Optional

import httpx
import jwt
from dotenv import load_dotenv
from fastapi import APIRouter, Depends, FastAPI, Header, HTTPException, Request
from fastapi.responses import PlainTextResponse
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, EmailStr, Field
from starlette.middleware.cors import CORSMiddleware

from content import CONTENT, SEED_HOUSEHOLDS

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ.get("EMERGENT_EMAIL_KEY", "")
EMAIL_FROM_NAME = os.environ.get("EMAIL_FROM_NAME", "Sophie + Ken")
PREVIEW_PASSWORD = os.environ.get("PREVIEW_PASSWORD", "soph")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "skks-2027")
JWT_SECRET = os.environ.get("JWT_SECRET", "skks-jwt-secret")
SITE_URL = os.environ.get("SITE_URL", "").rstrip("/")

app = FastAPI()
api_router = APIRouter(prefix="/api")
logger = logging.getLogger(__name__)

GENERIC_LOOKUP_FAIL = "We couldn't find that name. Check the spelling, try another member of your household, or contact us and we'll sort it out."


def now_iso():
    return datetime.now(timezone.utc).isoformat()


_hits = defaultdict(list)


def rate_limit(key: str, limit: int, window: int):
    t = time.time()
    _hits[key] = [x for x in _hits[key] if t - x < window]
    if len(_hits[key]) >= limit:
        raise HTTPException(status_code=429, detail="Too many requests. Please wait a moment and try again.")
    _hits[key].append(t)


# ---------- name normalization ----------

def normalize_name(name: str) -> str:
    """Trim, lowercase, collapse spaces, unify apostrophes, keep hyphens."""
    name = unicodedata.normalize("NFKC", name or "")
    name = name.replace("’", "'").replace("‘", "'").replace("`", "'").replace("ʼ", "'")
    name = name.lower().strip()
    name = re.sub(r"[^a-z0-9'\- ]", "", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name


def strip_punct(name: str) -> str:
    """Apostrophes and hyphens removed entirely, spaces collapsed."""
    return re.sub(r"\s+", " ", re.sub(r"['\-]", "", name)).strip()


def lookup_keys(first_name: str, last_name: str):
    full = normalize_name(f"{first_name} {last_name}")
    return full, strip_punct(full), re.sub(r"\s+", " ", full.replace("-", " ")).strip()


# ---------- models ----------

class PreviewUnlock(BaseModel):
    password: str


class GuestLookupIn(BaseModel):
    name: str = Field(default="", max_length=160)
    disambiguator: Optional[str] = Field(default=None, max_length=160)
    candidates_token: Optional[str] = None
    company: Optional[str] = ""  # honeypot
    t: Optional[int] = 0  # form render epoch ms (time-trap)


class GuestResponseIn(BaseModel):
    guest_id: str = Field(max_length=64)
    attending: bool
    meal_choice: Optional[str] = Field(default="", max_length=120)
    dietary_notes: Optional[str] = Field(default="", max_length=500)
    plus_one_name: Optional[str] = Field(default="", max_length=120)


class EventResponseIn(BaseModel):
    guest_id: str = Field(max_length=64)
    event_code: str = Field(max_length=40)
    attending: bool


class SubmitRsvpIn(BaseModel):
    token: str = Field(max_length=2048)
    email: EmailStr
    responses: List[GuestResponseIn]
    event_responses: List[EventResponseIn] = []
    accessibility_notes: Optional[str] = Field(default="", max_length=500)
    song_request: Optional[str] = Field(default="", max_length=200)
    message_to_couple: Optional[str] = Field(default="", max_length=1000)
    company: Optional[str] = ""  # honeypot
    t: Optional[int] = 0


class AdminLogin(BaseModel):
    password: str


class MemberIn(BaseModel):
    first_name: str = Field(max_length=80)
    last_name: str = Field(max_length=80)
    plus_one_allowed: bool = False


class HouseholdIn(BaseModel):
    name: str = Field(max_length=160)
    members: List[MemberIn]


class ImportPreviewIn(BaseModel):
    source: str = "text"  # "text" | "url" | "xlsx"
    data: str = Field(max_length=30_000_000)  # xlsx arrives base64-encoded


class ImportCommitIn(BaseModel):
    token: str = Field(max_length=128)


class SettingsIn(BaseModel):
    deadline: Optional[str] = None
    meal_options_enabled: bool = False
    meal_options: List[str] = []


class AdminResponseIn(BaseModel):
    email: Optional[str] = Field(default="", max_length=160)
    responses: List[GuestResponseIn]
    accessibility_notes: Optional[str] = Field(default="", max_length=500)
    song_request: Optional[str] = Field(default="", max_length=200)
    message_to_couple: Optional[str] = Field(default="", max_length=1000)


# ---------- auth ----------

def require_admin(authorization: str = Header(None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Unauthorized")
    try:
        jwt.decode(authorization[7:], JWT_SECRET, algorithms=["HS256"])
    except Exception:
        raise HTTPException(status_code=401, detail="Unauthorized")


def sign_token(payload: dict, minutes: int = None, days: int = None) -> str:
    exp = datetime.now(timezone.utc) + (timedelta(days=days) if days else timedelta(minutes=minutes or 30))
    return jwt.encode({**payload, "exp": exp}, JWT_SECRET, algorithm="HS256")


def read_token(token: str, scopes) -> dict:
    try:
        data = jwt.decode(token, JWT_SECRET, algorithms=["HS256"])
    except Exception:
        raise HTTPException(status_code=401, detail="That link has expired. Please look up your invitation again.")
    if data.get("scope") not in scopes:
        raise HTTPException(status_code=401, detail="That link has expired. Please look up your invitation again.")
    return data


# ---------- settings / deadline ----------

DEFAULT_SETTINGS = {"id": "rsvp", "deadline": None, "meal_options_enabled": False, "meal_options": []}


async def get_settings() -> dict:
    s = await db.settings.find_one({"id": "rsvp"}, {"_id": 0})
    return {**DEFAULT_SETTINGS, **(s or {})}


def deadline_passed(deadline: Optional[str]) -> bool:
    if not deadline:
        return False
    try:
        return datetime.now(timezone.utc) > datetime.fromisoformat(deadline)
    except ValueError:
        return False


async def assert_rsvp_open():
    s = await get_settings()
    if deadline_passed(s.get("deadline")):
        raise HTTPException(
            status_code=403,
            detail="RSVPs are now closed. If your plans have changed, please contact us directly.",
        )


# ---------- helpers ----------

def event_public(e, attendance=None):
    """Event shape safe for an invited guest. Null fields stay null — never invent details."""
    return {
        "code": e["event_code"],
        "name": e.get("event_name") or e["event_code"],
        "date": e.get("date"),
        "start_time": e.get("start_time"),
        "end_time": e.get("end_time"),
        "location": e.get("location"),
        "address": e.get("address"),
        "attendance": attendance,
    }


async def member_events(member_ids):
    """Map guest_id -> [invited secondary events]. WEDDING is implied and excluded."""
    if not member_ids:
        return {}
    invites = await db.event_invites.find(
        {"guest_id": {"$in": member_ids}, "invited": True, "event_code": {"$ne": "WEDDING"}},
        {"_id": 0},
    ).to_list(500)
    if not invites:
        return {}
    events = {e["event_code"]: e for e in await db.events.find({}, {"_id": 0}).to_list(50)}
    out = defaultdict(list)
    for inv in invites:
        e = events.get(inv["event_code"])
        if e and e.get("rsvp_enabled", True):
            out[inv["guest_id"]].append(event_public(e, inv.get("attendance")))
    return out


def member_public(g):
    return {
        "id": g["id"],
        "first_name": g["first_name"],
        "last_name": g["last_name"],
        "plus_one_allowed": g.get("plus_one_allowed", False),
    }


async def party_payload(party: dict, include_existing: bool = True):
    members = await db.guests.find(
        {"party_id": party["id"], "is_plus_one": {"$ne": True}}, {"_id": 0}
    ).to_list(50)
    events_by_guest = await member_events([m["id"] for m in members])
    members_out = []
    for m in members:
        mp = member_public(m)
        evs = events_by_guest.get(m["id"])
        if evs:  # never leak event data to guests with no invitation — omit the key entirely
            mp["events"] = evs
        members_out.append(mp)
    existing = None
    if include_existing:
        detail = await db.rsvp_details.find_one({"party_id": party["id"]}, {"_id": 0})
        responded = [m for m in members if m.get("attendance") in ("yes", "no")]
        if party.get("submitted_at") or detail or responded:
            existing = {
                "email": party.get("contact_email", ""),
                "responses": [
                    {
                        "guest_id": m["id"],
                        "attending": m.get("attendance") == "yes",
                        "meal_choice": m.get("meal_choice", ""),
                        "dietary_notes": m.get("dietary_notes", ""),
                        "plus_one_name": m.get("plus_one_name", ""),
                    }
                    for m in members
                ],
                "event_responses": [
                    {"guest_id": gid, "event_code": ev["code"], "attending": ev["attendance"] == "yes"}
                    for gid, evs in events_by_guest.items()
                    for ev in evs
                    if ev.get("attendance") in ("yes", "no")
                ],
                "accessibility_notes": (detail or {}).get("accessibility_notes", ""),
                "song_request": (detail or {}).get("song_request", ""),
                "message_to_couple": (detail or {}).get("message_to_couple", ""),
            }
    return {
        "party": {"display_name": party["display_name"]},
        "members": members_out,
        "already_responded": existing is not None,
        "existing": existing,
    }


async def send_email(recipient: str, subject: str, html: str):
    payload = {"to": [recipient], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
    try:
        async with httpx.AsyncClient(timeout=30) as client_http:
            resp = await client_http.post(
                f"{EMAIL_BASE_URL}/api/v1/email/send",
                headers={"X-Email-Key": EMAIL_KEY},
                json=payload,
            )
        if resp.status_code >= 400:
            import json as _json
            logger.error(f"Email service {resp.status_code}: {resp.text[:300]} | payload={_json.dumps({k: (v[:80] if isinstance(v, str) else v) for k, v in payload.items()})}")
        resp.raise_for_status()
        return True
    except Exception as e:
        logger.error(f"Email send error: {e}")
        return False


def esc(s) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_confirmation_email(party_name, responses, attending_any, details, update_url, settings, event_responses=None):
    rows = ""
    for r in responses:
        extras = []
        if r.get("plus_one_name"):
            extras.append(f"brings {esc(r['plus_one_name'])}")
        if settings.get("meal_options_enabled") and r.get("meal_choice"):
            extras.append(esc(r["meal_choice"]))
        if r.get("dietary_notes"):
            extras.append(f"kitchen note: {esc(r['dietary_notes'])}")
        extra_html = (
            f"<br><span style='font-size:12px;color:#595959;'>{' &middot; '.join(extras)}</span>" if extras else ""
        )
        rows += (
            f"<tr>"
            f"<td style='padding:8px 16px 8px 0;font-family:Georgia,serif;font-size:15px;color:#1A1A1A;border-bottom:1px solid #e3ded4;'>"
            f"{esc(r['guest_name'])}{extra_html}</td>"
            f"<td style='padding:8px 0;font-family:Georgia,serif;font-size:13px;letter-spacing:0.08em;text-transform:uppercase;"
            f"color:{'#4A5D4E' if r['attending'] else '#731F17'};border-bottom:1px solid #e3ded4;vertical-align:top;'>"
            f"{'Attending' if r['attending'] else 'Regrets'}</td>"
            f"</tr>"
        )
    extra_rows = ""
    if event_responses:
        rd_rows = "".join(
            f"<tr>"
            f"<td style='padding:6px 16px 6px 0;font-family:Georgia,serif;font-size:14px;color:#1A1A1A;border-bottom:1px solid #e3ded4;'>{esc(r['guest_name'])}</td>"
            f"<td style='padding:6px 0;font-family:Georgia,serif;font-size:12px;letter-spacing:0.08em;text-transform:uppercase;"
            f"color:{'#4A5D4E' if r['attending'] else '#731F17'};border-bottom:1px solid #e3ded4;'>"
            f"{'Attending' if r['attending'] else 'Regrets'}</td>"
            f"</tr>"
            for r in event_responses
        )
        extra_rows += (
            "<tr><td style='padding-top:24px;font-family:Courier New,monospace;font-size:11px;letter-spacing:0.2em;color:#595959;'>"
            "THE NIGHT BEFORE &nbsp;&middot;&nbsp; REHEARSAL DINNER</td></tr>"
            f"<tr><td><table role='presentation' width='100%' cellpadding='0' cellspacing='0' style='margin-top:8px;'>{rd_rows}</table></td></tr>"
            "<tr><td style='padding-top:6px;font-family:Georgia,serif;font-style:italic;font-size:12px;color:#595959;'>Details to follow.</td></tr>"
        )
    if details.get("accessibility_notes"):
        extra_rows += f"<tr><td style='padding-top:16px;font-family:Georgia,serif;font-size:13px;color:#595959;'>Comfort notes: {esc(details['accessibility_notes'])}</td></tr>"
    if details.get("song_request"):
        extra_rows += f"<tr><td style='padding-top:8px;font-family:Georgia,serif;font-size:13px;color:#595959;'>Song request: &ldquo;{esc(details['song_request'])}&rdquo; (a request, not a legally binding agreement)</td></tr>"
    if details.get("message_to_couple"):
        extra_rows += f"<tr><td style='padding-top:8px;font-family:Georgia,serif;font-size:13px;color:#595959;'>Your note: &ldquo;{esc(details['message_to_couple'])}&rdquo;</td></tr>"

    headline = "YOU&rsquo;RE IN." if attending_any else "WE&rsquo;LL MISS YOU."
    sub = (
        "We&rsquo;ll see you in New York on June 5."
        if attending_any
        else "Thank you for letting us know. We&rsquo;re lucky to have you in our lives, wherever you are that night."
    )
    update_html = ""
    if update_url:
        update_html = f"""
      <tr><td align="center" style="padding-top:28px;">
        <a href="{update_url}" style="font-family:'Courier New',monospace;font-size:11px;letter-spacing:0.2em;color:#731F17;text-decoration:underline;">PLANS CHANGE? VIEW OR UPDATE YOUR RSVP</a>
      </td></tr>"""
    return f"""
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#F7F5F0;padding:32px 16px;">
  <tr><td align="center">
    <table role="presentation" width="560" cellpadding="0" cellspacing="0" style="background:#F7F5F0;border:1px solid #d8d2c4;padding:48px 40px;">
      <tr><td align="center" style="font-family:'Courier New',monospace;font-size:11px;letter-spacing:0.25em;color:#595959;padding-bottom:24px;">
        VOL. I &nbsp;&middot;&nbsp; NO. 1 &nbsp;&middot;&nbsp; THE WEDDING ISSUE
      </td></tr>
      <tr><td align="center" style="font-family:Georgia,serif;font-size:30px;color:#1A1A1A;padding-bottom:8px;">
        Sophie <span style="color:#731F17;">&amp;</span> Ken
      </td></tr>
      <tr><td align="center" style="font-family:'Courier New',monospace;font-size:11px;letter-spacing:0.2em;color:#595959;padding-bottom:32px;">
        SATURDAY, JUNE 5, 2027 &nbsp;&middot;&nbsp; NEW YORK ATHLETIC CLUB
      </td></tr>
      <tr><td align="center" style="font-family:Georgia,serif;font-size:24px;color:#731F17;padding-bottom:12px;">{headline}</td></tr>
      <tr><td align="center" style="font-family:Georgia,serif;font-style:italic;font-size:15px;color:#595959;padding-bottom:32px;">{sub}</td></tr>
      <tr><td>
        <table role="presentation" width="100%" cellpadding="0" cellspacing="0">{rows}</table>
        <table role="presentation" width="100%" cellpadding="0" cellspacing="0">{extra_rows}</table>
      </td></tr>
      <tr><td align="center" style="padding-top:36px;">
        <span style="font-family:'Courier New',monospace;font-size:11px;letter-spacing:0.2em;color:#731F17;border:1px solid #731F17;display:inline-block;padding:12px 24px;">
        JUN 05 2027 &nbsp;&middot;&nbsp; 180 CENTRAL PARK SOUTH, NEW YORK
        </span>
      </td></tr>
      <tr><td align="center" style="padding-top:24px;font-family:Georgia,serif;font-size:12px;color:#595959;">
        Doors 5:30 PM &nbsp;&middot;&nbsp; Ceremony 6:00 PM
      </td></tr>{update_html}
    </table>
  </td></tr>
</table>"""


# ---------- public routes ----------

@api_router.get("/")
async def root():
    return {"message": "Sophie + Ken — June 5, 2027"}


@api_router.get("/content")
async def get_content():
    return CONTENT


@api_router.post("/preview/unlock")
async def preview_unlock(body: PreviewUnlock, request: Request):
    rate_limit(f"preview:{request.client.host}", 10, 300)
    if body.password != PREVIEW_PASSWORD:
        raise HTTPException(status_code=401, detail="Incorrect password")
    return {"ok": True}


@api_router.get("/rsvp/status")
async def rsvp_status():
    s = await get_settings()
    out = {
        "closed": deadline_passed(s.get("deadline")),
        "deadline": s.get("deadline"),
        "meal_options_enabled": s.get("meal_options_enabled", False),
    }
    if out["meal_options_enabled"]:
        out["meal_options"] = s.get("meal_options", [])
    return out


@api_router.post("/guest-lookup")
async def guest_lookup(body: GuestLookupIn, request: Request):
    rate_limit(f"lookup:{request.client.host}", 20, 300)
    if body.company:  # honeypot — behave like a miss
        return {"status": "not_found"}
    if body.t:
        elapsed_ms = int(time.time() * 1000) - body.t
        if 0 <= elapsed_ms < 1200:  # impossibly fast for a human
            return {"status": "not_found"}

    # --- disambiguation round ---
    if body.candidates_token and body.disambiguator:
        data = read_token(body.candidates_token, ["candidates"])
        candidate_ids = data.get("party_ids", [])
        d = body.disambiguator.strip().lower()
        draw = normalize_name(body.disambiguator)
        dn, da = draw, strip_punct(draw)
        ds = re.sub(r"\s+", " ", draw.replace("-", " ")).strip()
        for pid in candidate_ids:
            party = await db.parties.find_one({"id": pid, "status": "live"}, {"_id": 0})
            if not party:
                continue
            if party.get("contact_email") and party["contact_email"].lower() == d:
                payload = await party_payload(party)
                return {"status": "found", "token": sign_token({"scope": "party", "party_id": pid}, minutes=60), **payload}
            member = await db.guests.find_one(
                {"party_id": pid, "is_plus_one": {"$ne": True},
                 "$or": [{"lookup_name": dn}, {"lookup_alt": da}, {"lookup_spaced": ds}]},
                {"_id": 0},
            )
            if member:
                payload = await party_payload(party)
                return {"status": "found", "token": sign_token({"scope": "party", "party_id": pid}, minutes=60), **payload}
        return {"status": "not_found"}

    # --- initial name lookup ---
    norm = normalize_name(body.name)
    if not norm:
        return {"status": "not_found"}
    alt = strip_punct(norm)
    spaced = re.sub(r"\s+", " ", norm.replace("-", " ")).strip()
    matches = await db.guests.find(
        {"is_plus_one": {"$ne": True},
         "$or": [{"lookup_name": norm}, {"lookup_alt": alt}, {"lookup_spaced": spaced}]},
        {"_id": 0, "party_id": 1},
    ).to_list(100)
    party_ids = list({m["party_id"] for m in matches})
    live_parties = []
    for pid in party_ids:
        p = await db.parties.find_one({"id": pid, "status": "live"}, {"_id": 0})
        if p:
            live_parties.append(p)

    if not live_parties:
        return {"status": "not_found"}
    if len(live_parties) > 1:
        method = "email" if any(p.get("contact_email") for p in live_parties) else "member"
        return {
            "status": "multiple",
            "method": method,
            "candidates_token": sign_token({"scope": "candidates", "party_ids": [p["id"] for p in live_parties]}, minutes=15),
        }

    party = live_parties[0]
    payload = await party_payload(party)
    return {"status": "found", "token": sign_token({"scope": "party", "party_id": party["id"]}, minutes=60), **payload}


@api_router.get("/rsvp/update/{token}")
async def rsvp_update_load(token: str, request: Request):
    rate_limit(f"update:{request.client.host}", 30, 300)
    await assert_rsvp_open()
    data = read_token(token, ["update"])
    party = await db.parties.find_one({"id": data["party_id"], "status": "live"}, {"_id": 0})
    if not party:
        raise HTTPException(status_code=404, detail=GENERIC_LOOKUP_FAIL)
    payload = await party_payload(party)
    return {"status": "found", "token": token, **payload}


@api_router.post("/rsvp/submit")
async def rsvp_submit(body: SubmitRsvpIn, request: Request):
    rate_limit(f"submit:{request.client.host}", 10, 300)
    if body.company:
        raise HTTPException(status_code=400, detail="Something went wrong. Please try again.")
    if body.t:
        elapsed_ms = int(time.time() * 1000) - body.t
        if 0 <= elapsed_ms < 1000:
            raise HTTPException(status_code=400, detail="Something went wrong. Please try again.")
    await assert_rsvp_open()

    data = read_token(body.token, ["party", "update"])
    party_id = data["party_id"]  # never trust a client-supplied party id
    party = await db.parties.find_one({"id": party_id, "status": "live"}, {"_id": 0})
    if not party:
        raise HTTPException(status_code=404, detail="Invitation not found.")

    members = await db.guests.find({"party_id": party_id, "is_plus_one": {"$ne": True}}, {"_id": 0}).to_list(50)
    member_map = {m["id"]: m for m in members}
    if not body.responses:
        raise HTTPException(status_code=400, detail="No responses provided.")
    if len(body.responses) > len(members):
        raise HTTPException(status_code=400, detail="More responses than invited guests.")

    settings = await get_settings()
    meal_enabled = settings.get("meal_options_enabled", False)
    allowed_meals = set(settings.get("meal_options", []))

    responses_out = []
    seen = set()
    for r in body.responses:
        if r.guest_id in seen:
            raise HTTPException(status_code=400, detail="Duplicate response for a guest.")
        seen.add(r.guest_id)
        m = member_map.get(r.guest_id)
        if not m:
            raise HTTPException(status_code=400, detail="Guest does not belong to this invitation.")
        if not isinstance(r.attending, bool):
            raise HTTPException(status_code=400, detail="Invalid attendance value.")
        if r.plus_one_name and not m.get("plus_one_allowed"):
            raise HTTPException(status_code=400, detail="This guest does not have plus-one access.")
        if r.meal_choice:
            if not meal_enabled or r.meal_choice not in allowed_meals:
                raise HTTPException(status_code=400, detail="Invalid meal selection.")
        responses_out.append({
            "guest_id": r.guest_id,
            "guest_name": f"{m['first_name']} {m['last_name']}",
            "attending": r.attending,
            "meal_choice": r.meal_choice.strip() if r.attending else "",
            "dietary_notes": r.dietary_notes.strip() if r.attending else "",
            "plus_one_name": r.plus_one_name.strip() if (r.attending and r.plus_one_name) else "",
        })

    attending_any = any(r["attending"] for r in responses_out)
    ts = now_iso()

    # --- secondary event responses (e.g. rehearsal dinner) ---
    # Eligibility is verified server-side: the guest must belong to this party
    # AND hold an invited event_invites record. No self-invitation possible.
    event_responses_out = []
    events = {e["event_code"]: e for e in await db.events.find({}, {"_id": 0}).to_list(50)}
    for er in body.event_responses:
        m = member_map.get(er.guest_id)
        if not m:
            raise HTTPException(status_code=400, detail="Guest does not belong to this invitation.")
        if er.event_code == "WEDDING":
            raise HTTPException(status_code=400, detail="Invalid event.")
        invite = await db.event_invites.find_one(
            {"guest_id": er.guest_id, "event_code": er.event_code, "invited": True}, {"_id": 0}
        )
        if not invite:
            raise HTTPException(status_code=400, detail="This guest is not invited to that event.")
        if not isinstance(er.attending, bool):
            raise HTTPException(status_code=400, detail="Invalid attendance value.")
        await db.event_invites.update_one(
            {"id": invite["id"]},
            {"$set": {"attendance": "yes" if er.attending else "no", "updated_at": ts}},
        )
        ev = events.get(er.event_code, {})
        event_responses_out.append({
            "guest_name": f"{m['first_name']} {m['last_name']}",
            "event_code": er.event_code,
            "event_name": ev.get("event_name") or er.event_code,
            "attending": er.attending,
        })

    for r in responses_out:
        await db.guests.update_one(
            {"id": r["guest_id"]},
            {"$set": {
                "attendance": "yes" if r["attending"] else "no",
                "meal_choice": r["meal_choice"],
                "dietary_notes": r["dietary_notes"],
                "plus_one_name": r["plus_one_name"],
                "updated_at": ts,
            }},
        )
    await db.parties.update_one(
        {"id": party_id},
        {"$set": {"contact_email": body.email, "updated_at": ts},
         "$setOnInsert": {"submitted_at": ts}},
    )
    if not party.get("submitted_at"):
        await db.parties.update_one({"id": party_id}, {"$set": {"submitted_at": ts}})
    await db.rsvp_details.update_one(
        {"party_id": party_id},
        {"$set": {
            "party_id": party_id,
            "accessibility_notes": body.accessibility_notes.strip(),
            "song_request": body.song_request.strip(),
            "message_to_couple": body.message_to_couple.strip(),
            "updated_at": ts,
        }},
        upsert=True,
    )

    update_url = ""
    if SITE_URL:
        update_url = f"{SITE_URL}/rsvp?token={sign_token({'scope': 'update', 'party_id': party_id}, days=150)}"

    details = {
        "accessibility_notes": body.accessibility_notes.strip(),
        "song_request": body.song_request.strip(),
        "message_to_couple": body.message_to_couple.strip(),
    }
    html = build_confirmation_email(party["display_name"], responses_out, attending_any, details, update_url, settings, event_responses_out)
    subject = "You're in — Sophie + Ken, June 5, 2027" if attending_any else "We'll miss you — Sophie + Ken"
    email_sent = await send_email(body.email, subject, html)

    return {"ok": True, "attending_any": attending_any, "email_sent": email_sent, "responses": responses_out, "event_responses": event_responses_out}


# ---------- admin routes ----------

@api_router.post("/admin/login")
async def admin_login(body: AdminLogin, request: Request):
    rate_limit(f"admin:{request.client.host}", 10, 300)
    if body.password != ADMIN_PASSWORD:
        raise HTTPException(status_code=401, detail="Incorrect password")
    token = jwt.encode(
        {"role": "admin", "exp": datetime.now(timezone.utc) + timedelta(hours=12)},
        JWT_SECRET,
        algorithm="HS256",
    )
    return {"token": token}


@api_router.get("/admin/overview", dependencies=[Depends(require_admin)])
async def admin_overview():
    parties = await db.parties.find({}, {"_id": 0}).to_list(10000)
    guests = await db.guests.find({"is_plus_one": {"$ne": True}}, {"_id": 0}).to_list(10000)
    responded_parties = [p for p in parties if p.get("submitted_at")]
    attending = sum(1 for g in guests if g.get("attendance") == "yes")
    declined = sum(1 for g in guests if g.get("attendance") == "no")
    responded_guests = attending + declined
    plus_ones = sum(1 for g in guests if g.get("attendance") == "yes" and g.get("plus_one_name"))
    dietary = sum(1 for g in guests if g.get("attendance") == "yes" and g.get("dietary_notes"))
    details = await db.rsvp_details.find({}, {"_id": 0}).to_list(10000)
    accessibility = sum(1 for d in details if d.get("accessibility_notes"))
    settings = await get_settings()
    meal_totals = {}
    if settings.get("meal_options_enabled"):
        allowed = set(settings.get("meal_options", []))
        for g in guests:
            if g.get("attendance") == "yes" and g.get("meal_choice") in allowed:
                meal_totals[g["meal_choice"]] = meal_totals.get(g["meal_choice"], 0) + 1
    stamps = [p.get("updated_at") for p in parties if p.get("updated_at")] + [d.get("updated_at") for d in details if d.get("updated_at")]
    rd_invites = await db.event_invites.find({"event_code": "REHEARSAL_DINNER", "invited": True}, {"_id": 0}).to_list(10000)
    rd_attending = sum(1 for i in rd_invites if i.get("attendance") == "yes")
    rd_declined = sum(1 for i in rd_invites if i.get("attendance") == "no")
    return {
        "parties": len(parties),
        "total_invited": len(guests),
        "parties_responded": len(responded_parties),
        "attending": attending,
        "declined": declined,
        "outstanding": len(guests) - responded_guests,
        "plus_ones": plus_ones,
        "dietary_notes": dietary,
        "accessibility_notes": accessibility,
        "meal_totals": meal_totals,
        "last_updated": max(stamps) if stamps else None,
        "rehearsal_dinner": {
            "invited": len(rd_invites),
            "responded": rd_attending + rd_declined,
            "attending": rd_attending,
            "declined": rd_declined,
            "outstanding": len(rd_invites) - rd_attending - rd_declined,
        },
    }


@api_router.get("/admin/parties", dependencies=[Depends(require_admin)])
async def admin_parties(search: str = "", filter: str = "all"):
    parties = await db.parties.find({}, {"_id": 0}).to_list(10000)
    guests = await db.guests.find({"is_plus_one": {"$ne": True}}, {"_id": 0}).to_list(10000)
    details = {d["party_id"]: d for d in await db.rsvp_details.find({}, {"_id": 0}).to_list(10000)}
    by_party = defaultdict(list)
    for g in guests:
        by_party[g["party_id"]].append(g)

    q = normalize_name(search)
    out = []
    for p in parties:
        members = by_party[p["id"]]
        responded = bool(p.get("submitted_at"))
        attending_any = any(m.get("attendance") == "yes" for m in members)
        if filter == "responded" and not responded:
            continue
        if filter == "outstanding" and responded:
            continue
        if filter == "attending" and not (responded and attending_any):
            continue
        if filter == "declined" and not (responded and not attending_any):
            continue
        if q:
            haystack = [normalize_name(p["display_name"])] + [
                normalize_name(f"{m['first_name']} {m['last_name']}") for m in members
            ]
            if not any(q in h for h in haystack):
                continue
        out.append({
            "id": p["id"],
            "display_name": p["display_name"],
            "contact_email": p.get("contact_email", ""),
            "status": p.get("status", "live"),
            "responded": responded,
            "attending_any": attending_any if responded else None,
            "submitted_at": p.get("submitted_at"),
            "updated_at": p.get("updated_at"),
            "members": [
                {
                    **member_public(m),
                    "attendance": m.get("attendance"),
                    "meal_choice": m.get("meal_choice", ""),
                    "dietary_notes": m.get("dietary_notes", ""),
                    "plus_one_name": m.get("plus_one_name", ""),
                }
                for m in members
            ],
            "details": details.get(p["id"], {}),
        })
    out.sort(key=lambda x: x["display_name"].lower())
    return out


@api_router.put("/admin/parties/{pid}/response", dependencies=[Depends(require_admin)])
async def admin_update_response(pid: str, body: AdminResponseIn):
    party = await db.parties.find_one({"id": pid}, {"_id": 0})
    if not party:
        raise HTTPException(status_code=404, detail="Party not found")
    members = await db.guests.find({"party_id": pid, "is_plus_one": {"$ne": True}}, {"_id": 0}).to_list(50)
    member_map = {m["id"]: m for m in members}
    ts = now_iso()
    for r in body.responses:
        m = member_map.get(r.guest_id)
        if not m:
            raise HTTPException(status_code=400, detail="Guest does not belong to this party.")
        if r.plus_one_name and not m.get("plus_one_allowed"):
            raise HTTPException(status_code=400, detail="This guest does not have plus-one access.")
        await db.guests.update_one(
            {"id": r.guest_id},
            {"$set": {
                "attendance": "yes" if r.attending else "no",
                "meal_choice": r.meal_choice.strip() if r.attending else "",
                "dietary_notes": r.dietary_notes.strip() if r.attending else "",
                "plus_one_name": r.plus_one_name.strip() if (r.attending and r.plus_one_name) else "",
                "updated_at": ts,
            }},
        )
    await db.parties.update_one(
        {"id": pid},
        {"$set": {"contact_email": body.email or party.get("contact_email", ""), "updated_at": ts, "submitted_at": party.get("submitted_at") or ts}},
    )
    await db.rsvp_details.update_one(
        {"party_id": pid},
        {"$set": {
            "party_id": pid,
            "accessibility_notes": body.accessibility_notes.strip(),
            "song_request": body.song_request.strip(),
            "message_to_couple": body.message_to_couple.strip(),
            "updated_at": ts,
        }},
        upsert=True,
    )
    return {"ok": True}


@api_router.get("/admin/settings", dependencies=[Depends(require_admin)])
async def admin_get_settings():
    return await get_settings()


@api_router.put("/admin/settings", dependencies=[Depends(require_admin)])
async def admin_put_settings(body: SettingsIn):
    deadline = body.deadline or None
    if deadline:
        try:
            datetime.fromisoformat(deadline)
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid deadline format.")
    await db.settings.update_one(
        {"id": "rsvp"},
        {"$set": {
            "id": "rsvp",
            "deadline": deadline,
            "meal_options_enabled": body.meal_options_enabled,
            "meal_options": [m.strip() for m in body.meal_options if m.strip()][:20],
        }},
        upsert=True,
    )
    return {"ok": True}


# ---------- import (CSV / Google Sheet / xlsx workbook) ----------

_import_cache = {}


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, bool):
        return "yes" if v else "no"
    return str(v).strip()


def truthy(v):
    return cell_str(v).lower() in ("yes", "true", "1", "y")


def split_full_name(full: str):
    parts = full.strip().split()
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return " ".join(parts[:-1]), parts[-1]


def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [cell_str(h).lower() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({headers[i]: r[i] for i in range(min(len(headers), len(r))) if headers[i]})
    return headers, out


def parse_guest_workbook(b64_data: str):
    """Parse the SK27 import workbook: Events, Parties, Guests, Event_Invites sheets."""
    import base64
    import openpyxl

    try:
        raw = base64.b64decode(b64_data)
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read that workbook. Upload the .xlsx file itself.")

    events, parties, guests, invites, errors = [], [], [], [], []

    if "Events" in wb.sheetnames:
        for r in sheet_rows(wb["Events"])[1]:
            code = cell_str(r.get("event_code")).upper()
            if not code:
                continue
            events.append({
                "event_code": code,
                "event_name": cell_str(r.get("event_name")) or code,
                "date": cell_str(r.get("date")) or None,
                "start_time": cell_str(r.get("start_time")) or None,
                "end_time": cell_str(r.get("end_time")) or None,
                "location": cell_str(r.get("location")) or None,
                "address": cell_str(r.get("address")) or None,
                "rsvp_enabled": truthy(r.get("rsvp_enabled")) if cell_str(r.get("rsvp_enabled")) else True,
                "visibility": cell_str(r.get("visibility")) or "invited",
                "details_status": cell_str(r.get("details_status")) or None,
                "notes": cell_str(r.get("notes")) or None,
            })
    else:
        errors.append("No Events sheet found.")

    for r in sheet_rows(wb["Parties"])[1] if "Parties" in wb.sheetnames else []:
        epid = cell_str(r.get("external_party_id"))
        name = cell_str(r.get("display_name"))
        if not epid or not name:
            errors.append("A Parties row is missing external_party_id or display_name — skipped.")
            continue
        parties.append({
            "external_party_id": epid,
            "display_name": name,
            "contact_email": cell_str(r.get("contact_email")),
            "status": cell_str(r.get("status")) or "live",
        })

    for r in sheet_rows(wb["Guests"])[1] if "Guests" in wb.sheetnames else []:
        egid = cell_str(r.get("external_guest_id"))
        epid = cell_str(r.get("external_party_id"))
        full = cell_str(r.get("full_name")) or cell_str(r.get("source_name"))
        if not egid or not epid or not full:
            errors.append("A Guests row is missing external_guest_id, external_party_id, or full_name — skipped.")
            continue
        first, last = split_full_name(full)
        guests.append({
            "external_guest_id": egid,
            "external_party_id": epid,
            "first_name": first,
            "last_name": last,
            "plus_one_allowed": truthy(r.get("plus_one_allowed")),
            "is_plus_one": truthy(r.get("is_plus_one")),
            "status": cell_str(r.get("status")) or "live",
        })

    for r in sheet_rows(wb["Event_Invites"])[1] if "Event_Invites" in wb.sheetnames else []:
        eiid = cell_str(r.get("external_invite_id"))
        code = cell_str(r.get("event_code")).upper()
        egid = cell_str(r.get("external_guest_id"))
        if not eiid or not code or not egid:
            errors.append("An Event_Invites row is missing external_invite_id, event_code, or external_guest_id — skipped.")
            continue
        invites.append({
            "external_invite_id": eiid,
            "event_code": code,
            "external_guest_id": egid,
            "invited": truthy(r.get("invited")) if cell_str(r.get("invited")) else True,
        })

    return {"events": events, "parties": parties, "guests": guests, "event_invites": invites}, errors


def parse_guest_csv(text: str):
    reader = csv.DictReader(io.StringIO(text.strip()))
    if not reader.fieldnames:
        return [], ["The file appears to be empty."]
    headers = {h.strip().lower(): h for h in reader.fieldnames}
    def col(*names):
        for n in names:
            if n in headers:
                return headers[n]
        return None
    c_party = col("party", "household", "party_name", "household_name")
    c_first = col("first_name", "first")
    c_last = col("last_name", "last")
    c_email = col("email", "contact_email")
    c_plus = col("plus_one", "plus-one", "plusone", "plus_one_allowed")
    if not (c_party and c_first and c_last):
        return [], ["Missing required columns: party, first_name, last_name."]
    rows, errors = [], []
    for i, raw in enumerate(reader, start=2):
        party = (raw.get(c_party) or "").strip()
        first = (raw.get(c_first) or "").strip()
        last = (raw.get(c_last) or "").strip()
        email = (raw.get(c_email) or "").strip() if c_email else ""
        plus = (raw.get(c_plus) or "").strip().lower() in ("yes", "true", "1", "y") if c_plus else False
        if not party or not first or not last:
            errors.append(f"Row {i}: missing party, first name, or last name — skipped.")
            continue
        rows.append({"party": party, "first_name": first, "last_name": last, "email": email, "plus_one_allowed": plus})
    return rows, errors


@api_router.post("/admin/import/preview", dependencies=[Depends(require_admin)])
async def admin_import_preview(body: ImportPreviewIn):
    # --- xlsx workbook path (Parties / Guests / Events / Event_Invites) ---
    if body.source == "xlsx":
        parsed, errors = parse_guest_workbook(body.data)
        if not parsed["parties"] and not parsed["guests"]:
            raise HTTPException(status_code=400, detail=errors[0] if errors else "No valid rows found in that workbook.")
        token = secrets.token_urlsafe(16)
        _import_cache[token] = {"kind": "xlsx", "data": parsed, "expires": time.time() + 900}
        rd_count = sum(1 for i in parsed["event_invites"] if i["event_code"] == "REHEARSAL_DINNER" and i["invited"])
        return {
            "token": token,
            "kind": "xlsx",
            "events": len(parsed["events"]),
            "parties": len(parsed["parties"]),
            "guests": len(parsed["guests"]),
            "event_invites": len(parsed["event_invites"]),
            "rehearsal_dinner_invites": rd_count,
            "duplicates": [],
            "errors": errors,
            "sample": [
                {"party": p["display_name"], "members": []}
                for p in parsed["parties"][:8]
            ],
        }

    text = body.data
    if body.source == "url":
        url = body.data.strip()
        if not url.startswith("https://"):
            raise HTTPException(status_code=400, detail="Only https:// sheet URLs are allowed.")
        try:
            async with httpx.AsyncClient(timeout=30, follow_redirects=True) as http:
                resp = await http.get(url)
            resp.raise_for_status()
            text = resp.text
        except Exception:
            raise HTTPException(status_code=400, detail="Could not fetch that sheet. Make sure it is published to the web as CSV.")
    rows, errors = parse_guest_csv(text)
    if not rows:
        raise HTTPException(status_code=400, detail=errors[0] if errors else "No valid rows found.")

    existing = await db.guests.find({"is_plus_one": {"$ne": True}}, {"_id": 0, "lookup_name": 1, "lookup_alt": 1}).to_list(100000)
    existing_keys = {(g.get("lookup_name"), g.get("lookup_alt")) for g in existing}

    parties = defaultdict(list)
    for r in rows:
        parties[r["party"]].append(r)

    duplicates, clean_rows, seen_in_file = [], [], set()
    for r in rows:
        ln, la, ls = lookup_keys(r["first_name"], r["last_name"])
        file_key = (normalize_name(r["party"]), ln)
        if file_key in seen_in_file:
            duplicates.append(f"{r['first_name']} {r['last_name']} ({r['party']}) — twice in file")
            continue
        seen_in_file.add(file_key)
        if (ln, la) in existing_keys:
            duplicates.append(f"{r['first_name']} {r['last_name']} ({r['party']}) — already on the list")
            continue
        clean_rows.append(r)

    token = secrets.token_urlsafe(16)
    _import_cache[token] = {"rows": clean_rows, "expires": time.time() + 900}
    clean_parties = defaultdict(list)
    for r in clean_rows:
        clean_parties[r["party"]].append(r)
    return {
        "token": token,
        "parties": len(clean_parties),
        "guests": len(clean_rows),
        "duplicates": duplicates,
        "errors": errors,
        "sample": [
            {"party": p, "members": [f"{m['first_name']} {m['last_name']}" for m in ms]}
            for p, ms in list(clean_parties.items())[:8]
        ],
    }


@api_router.post("/admin/import/commit", dependencies=[Depends(require_admin)])
async def admin_import_commit(body: ImportCommitIn):
    cached = _import_cache.pop(body.token, None)
    if not cached or cached["expires"] < time.time():
        raise HTTPException(status_code=400, detail="That import preview has expired. Run it again.")

    # --- xlsx workbook commit: upsert by immutable external IDs, never by name ---
    if cached.get("kind") == "xlsx":
        parsed = cached["data"]
        ts = now_iso()
        for e in parsed["events"]:
            await db.events.update_one(
                {"event_code": e["event_code"]},
                {"$set": {**e, "updated_at": ts}, "$setOnInsert": {"id": str(uuid.uuid4()), "created_at": ts}},
                upsert=True,
            )
        party_id_map = {}
        for p in parsed["parties"]:
            existing = await db.parties.find_one({"external_party_id": p["external_party_id"]}, {"_id": 0})
            if existing:
                party_id_map[p["external_party_id"]] = existing["id"]
                await db.parties.update_one(
                    {"id": existing["id"]},
                    {"$set": {"display_name": p["display_name"], "contact_email": p["contact_email"] or existing.get("contact_email", ""), "status": p["status"], "updated_at": ts}},
                )
            else:
                pid = str(uuid.uuid4())
                party_id_map[p["external_party_id"]] = pid
                await db.parties.insert_one({
                    "id": pid,
                    "external_party_id": p["external_party_id"],
                    "invite_code": secrets.token_urlsafe(6),
                    "display_name": p["display_name"],
                    "contact_email": p["contact_email"],
                    "status": p["status"],
                    "submitted_at": None,
                    "created_at": ts,
                    "updated_at": ts,
                })
        guest_id_map = {}
        for g in parsed["guests"]:
            pid = party_id_map.get(g["external_party_id"])
            if not pid:
                continue
            ln, la, ls = lookup_keys(g["first_name"], g["last_name"])
            doc = {
                "party_id": pid,
                "first_name": g["first_name"],
                "last_name": g["last_name"],
                "full_name": f"{g['first_name']} {g['last_name']}".strip(),
                "lookup_name": ln,
                "lookup_alt": la,
                "lookup_spaced": ls,
                "plus_one_allowed": g["plus_one_allowed"],
                "is_plus_one": g["is_plus_one"],
                "status": g["status"],
                "updated_at": ts,
            }
            existing = await db.guests.find_one({"external_guest_id": g["external_guest_id"]}, {"_id": 0})
            if existing:
                guest_id_map[g["external_guest_id"]] = existing["id"]
                await db.guests.update_one({"id": existing["id"]}, {"$set": doc})
            else:
                gid = str(uuid.uuid4())
                guest_id_map[g["external_guest_id"]] = gid
                await db.guests.insert_one({
                    "id": gid,
                    "external_guest_id": g["external_guest_id"],
                    **doc,
                    "attendance": None,
                    "meal_choice": "",
                    "dietary_notes": "",
                    "plus_one_name": "",
                    "created_at": ts,
                })
        n_invites = 0
        for inv in parsed["event_invites"]:
            gid = guest_id_map.get(inv["external_guest_id"])
            if not gid:
                continue
            await db.event_invites.update_one(
                {"external_invite_id": inv["external_invite_id"]},
                {"$set": {
                    "event_code": inv["event_code"],
                    "guest_id": gid,
                    "invited": inv["invited"],
                    "updated_at": ts,
                }, "$setOnInsert": {"id": str(uuid.uuid4()), "attendance": None, "created_at": ts}},
                upsert=True,
            )
            n_invites += 1
        return {
            "ok": True,
            "events": len(parsed["events"]),
            "parties": len(parsed["parties"]),
            "guests": len(parsed["guests"]),
            "event_invites": n_invites,
        }

    rows = cached["rows"]
    parties = defaultdict(list)
    for r in rows:
        parties[r["party"]].append(r)
    n_guests = 0
    for party_name, members in parties.items():
        pid = str(uuid.uuid4())
        email = next((m["email"] for m in members if m["email"]), "")
        await db.parties.insert_one({
            "id": pid,
            "invite_code": secrets.token_urlsafe(6),
            "display_name": party_name.strip(),
            "contact_email": email,
            "status": "live",
            "submitted_at": None,
            "created_at": now_iso(),
            "updated_at": now_iso(),
        })
        for m in members:
            ln, la, ls = lookup_keys(m["first_name"], m["last_name"])
            await db.guests.insert_one({
                "id": str(uuid.uuid4()),
                "party_id": pid,
                "first_name": m["first_name"].strip(),
                "last_name": m["last_name"].strip(),
                "full_name": f"{m['first_name'].strip()} {m['last_name'].strip()}",
                "lookup_name": ln,
                "lookup_alt": la,
            "lookup_spaced": ls,
                "plus_one_allowed": m["plus_one_allowed"],
                "is_plus_one": False,
                "attendance": None,
                "meal_choice": "",
                "dietary_notes": "",
                "plus_one_name": "",
                "created_at": now_iso(),
            })
            n_guests += 1
    return {"ok": True, "parties": len(parties), "guests": n_guests}


# ---------- legacy guest-list management (admin) ----------

@api_router.post("/admin/parties", dependencies=[Depends(require_admin)])
async def admin_create_party(body: HouseholdIn):
    if not body.members:
        raise HTTPException(status_code=400, detail="A party needs at least one guest.")
    pid = str(uuid.uuid4())
    await db.parties.insert_one({
        "id": pid,
        "invite_code": secrets.token_urlsafe(6),
        "display_name": body.name.strip(),
        "contact_email": "",
        "status": "live",
        "submitted_at": None,
        "created_at": now_iso(),
        "updated_at": now_iso(),
    })
    for m in body.members:
        ln, la, ls = lookup_keys(m.first_name, m.last_name)
        await db.guests.insert_one({
            "id": str(uuid.uuid4()),
            "party_id": pid,
            "first_name": m.first_name.strip(),
            "last_name": m.last_name.strip(),
            "full_name": f"{m.first_name.strip()} {m.last_name.strip()}",
            "lookup_name": ln,
            "lookup_alt": la,
            "lookup_spaced": ls,
            "plus_one_allowed": m.plus_one_allowed,
            "is_plus_one": False,
            "attendance": None,
            "meal_choice": "",
            "dietary_notes": "",
            "plus_one_name": "",
            "created_at": now_iso(),
        })
    return {"ok": True, "id": pid}


@api_router.delete("/admin/parties/{pid}", dependencies=[Depends(require_admin)])
async def admin_delete_party(pid: str):
    await db.parties.delete_one({"id": pid})
    await db.guests.delete_many({"party_id": pid})
    await db.rsvp_details.delete_many({"party_id": pid})
    return {"ok": True}


@api_router.post("/admin/parties/{pid}/members", dependencies=[Depends(require_admin)])
async def admin_add_member(pid: str, body: MemberIn):
    party = await db.parties.find_one({"id": pid}, {"_id": 0})
    if not party:
        raise HTTPException(status_code=404, detail="Party not found")
    gid = str(uuid.uuid4())
    ln, la, ls = lookup_keys(body.first_name, body.last_name)
    await db.guests.insert_one({
        "id": gid,
        "party_id": pid,
        "first_name": body.first_name.strip(),
        "last_name": body.last_name.strip(),
        "full_name": f"{body.first_name.strip()} {body.last_name.strip()}",
        "lookup_name": ln,
        "lookup_alt": la,
            "lookup_spaced": ls,
        "plus_one_allowed": body.plus_one_allowed,
        "is_plus_one": False,
        "attendance": None,
        "meal_choice": "",
        "dietary_notes": "",
        "plus_one_name": "",
        "created_at": now_iso(),
    })
    return {"ok": True, "id": gid}


@api_router.delete("/admin/members/{gid}", dependencies=[Depends(require_admin)])
async def admin_delete_member(gid: str):
    await db.guests.delete_one({"id": gid})
    return {"ok": True}


@api_router.get("/admin/export", dependencies=[Depends(require_admin)])
async def admin_export():
    parties = {p["id"]: p for p in await db.parties.find({}, {"_id": 0}).to_list(10000)}
    guests = await db.guests.find({"is_plus_one": {"$ne": True}}, {"_id": 0}).to_list(10000)
    details = {d["party_id"]: d for d in await db.rsvp_details.find({}, {"_id": 0}).to_list(10000)}
    rd = {i["guest_id"]: i.get("attendance") for i in await db.event_invites.find({"event_code": "REHEARSAL_DINNER", "invited": True}, {"_id": 0}).to_list(10000)}
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Party", "Guest", "Wedding", "Rehearsal Dinner", "Plus One", "Meal", "Dietary", "Accessibility", "Song", "Note", "Email", "Submitted", "Updated"])
    for g in sorted(guests, key=lambda x: (parties.get(x["party_id"], {}).get("display_name", ""), x["last_name"])):
        p = parties.get(g["party_id"], {})
        d = details.get(g["party_id"], {})
        rd_val = rd.get(g["id"])
        writer.writerow([
            p.get("display_name", ""),
            f"{g['first_name']} {g['last_name']}",
            "Yes" if g.get("attendance") == "yes" else ("No" if g.get("attendance") == "no" else ""),
            "Yes" if rd_val == "yes" else ("No" if rd_val == "no" else ("Invited" if g["id"] in rd else "")),
            g.get("plus_one_name", ""),
            g.get("meal_choice", ""),
            g.get("dietary_notes", ""),
            d.get("accessibility_notes", ""),
            d.get("song_request", ""),
            d.get("message_to_couple", ""),
            p.get("contact_email", ""),
            p.get("submitted_at") or "",
            p.get("updated_at") or "",
        ])
    return PlainTextResponse(
        buf.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=sophie-ken-rsvps.csv"},
    )


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")


# ---------- migration & seed ----------

def new_party_doc(name, email="", status="live", submitted_at=None, created=None):
    ts = created or now_iso()
    return {
        "id": str(uuid.uuid4()),
        "invite_code": secrets.token_urlsafe(6),
        "display_name": name,
        "contact_email": email,
        "status": status,
        "submitted_at": submitted_at,
        "created_at": ts,
        "updated_at": ts,
    }


def new_guest_doc(party_id, first, last, plus_one=False):
    ln, la, ls = lookup_keys(first, last)
    return {
        "id": str(uuid.uuid4()),
        "party_id": party_id,
        "first_name": first,
        "last_name": last,
        "full_name": f"{first} {last}",
        "lookup_name": ln,
        "lookup_alt": la,
            "lookup_spaced": ls,
        "plus_one_allowed": plus_one,
        "is_plus_one": False,
        "attendance": None,
        "meal_choice": "",
        "dietary_notes": "",
        "plus_one_name": "",
        "created_at": now_iso(),
    }


async def migrate_legacy():
    """One-time migration from households/rsvps to parties/guests/rsvp_details."""
    if await db.parties.count_documents({}) > 0:
        return
    households = await db.households.find({}, {"_id": 0}).to_list(10000)
    if not households:
        return
    logger.info("Migrating legacy households to parties…")
    for h in households:
        rsvp = await db.rsvps.find_one({"household_id": h["id"]}, {"_id": 0})
        party = {
            "id": h["id"],
            "invite_code": secrets.token_urlsafe(6),
            "display_name": h["name"],
            "contact_email": (rsvp or {}).get("email", ""),
            "status": "live",
            "submitted_at": (rsvp or {}).get("submitted_at"),
            "created_at": h.get("created_at", now_iso()),
            "updated_at": (rsvp or {}).get("updated_at", h.get("created_at", now_iso())),
        }
        await db.parties.insert_one(party)
        legacy_members = await db.guests.find({"household_id": h["id"]}, {"_id": 0}).to_list(50)
        resp_map = {r["guest_id"]: r for r in (rsvp or {}).get("responses", [])}
        access_notes = []
        for m in legacy_members:
            ln, la, ls = lookup_keys(m["first_name"], m["last_name"])
            r = resp_map.get(m["id"], {})
            if r.get("accessibility"):
                access_notes.append(r["accessibility"])
            await db.guests.update_one(
                {"id": m["id"]},
                {"$set": {
                    "party_id": h["id"],
                    "full_name": f"{m['first_name']} {m['last_name']}",
                    "lookup_name": ln,
                    "lookup_alt": la,
            "lookup_spaced": ls,
                    "is_plus_one": False,
                    "attendance": ("yes" if r.get("attending") else "no") if r else None,
                    "meal_choice": r.get("meal", ""),
                    "dietary_notes": r.get("dietary", ""),
                    "plus_one_name": r.get("plus_one_name", ""),
                    "updated_at": (rsvp or {}).get("updated_at", now_iso()),
                }},
            )
        if rsvp:
            await db.rsvp_details.insert_one({
                "party_id": h["id"],
                "accessibility_notes": "; ".join(access_notes),
                "song_request": rsvp.get("personality_answer", ""),
                "message_to_couple": rsvp.get("note", ""),
                "updated_at": rsvp.get("updated_at", now_iso()),
            })
    logger.info("Migration complete.")


async def seed_if_empty():
    if await db.parties.count_documents({}) > 0:
        return
    for h in SEED_HOUSEHOLDS:
        party = new_party_doc(h["name"])
        await db.parties.insert_one(party)
        for m in h["members"]:
            await db.guests.insert_one(new_guest_doc(party["id"], m["first_name"], m["last_name"], m["plus_one_allowed"]))
        logger.info("Seeded sample guest list")


async def seed_test_fixtures():
    """Idempotent extra fixtures covering lookup edge cases (apostrophes, hyphens, duplicates)."""
    if await db.settings.find_one({"id": "test_fixtures_v1"}):
        return
    fixtures = [
        ("The O'Brien Household", [("Saoirse", "O'Brien", False), ("Cillian", "O'Brien", False)], ""),
        ("The Doe-Smith Household", [("Jamie", "Doe-Smith", True)], ""),
        ("The Sample Cousins", [("Alex", "Sample", False), ("Jordan", "Sample", False)], ""),
    ]
    for name, members, email in fixtures:
        if await db.parties.find_one({"display_name": name}):
            continue
        party = new_party_doc(name, email=email)
        await db.parties.insert_one(party)
        for first, last, plus in members:
            await db.guests.insert_one(new_guest_doc(party["id"], first, last, plus))
    await db.settings.update_one({"id": "test_fixtures_v1"}, {"$set": {"id": "test_fixtures_v1"}}, upsert=True)
    logger.info("Seeded lookup edge-case fixtures")


async def backfill_lookup_keys():
    """Ensure every guest has the full set of normalized lookup keys."""
    async for g in db.guests.find({"lookup_spaced": {"$exists": False}}):
        ln, la, ls = lookup_keys(g["first_name"], g["last_name"])
        updates = {"lookup_name": ln, "lookup_alt": la, "lookup_spaced": ls}
        if not g.get("party_id") and g.get("household_id"):
            updates["party_id"] = g["household_id"]
        await db.guests.update_one({"id": g["id"]}, {"$set": updates})


@app.on_event("startup")
async def startup():
    await db.guests.create_index("lookup_name")
    await db.guests.create_index("lookup_alt")
    await db.guests.create_index("lookup_spaced")
    await db.guests.create_index("party_id")
    await db.parties.create_index("invite_code", unique=True)
    await db.parties.create_index("status")
    await db.events.create_index("event_code", unique=True)
    await db.event_invites.create_index("external_invite_id", unique=True, sparse=True)
    await db.event_invites.create_index([("guest_id", 1), ("event_code", 1)])
    await migrate_legacy()
    await seed_if_empty()
    await backfill_lookup_keys()
    await seed_test_fixtures()
    logger.info(f"Email key loaded: {'yes' if EMAIL_KEY else 'NO'} (len={len(EMAIL_KEY)}), site_url={SITE_URL}")


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
